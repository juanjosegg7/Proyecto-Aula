"""Convierte un libro de Excel en Frontend/datos/ejemplo-panaderia.js (datos de ejemplo de la app).

Uso: python Simulacion/exportar_ejemplo.py Frontend/datos/Base_Panaderia_Indicadores_2026-09.xlsx
Las fechas se guardan como número de serie de Excel, igual que las lee la aplicación.
"""
import datetime
import json
import sys

from openpyxl import load_workbook

ORIGEN = datetime.datetime(1899, 12, 30)


def valor(v):
    if isinstance(v, datetime.datetime):
        return round((v - ORIGEN).total_seconds() / 86400, 8)
    if isinstance(v, datetime.date):
        return (v - ORIGEN.date()).days
    return v


def main(ruta, salida="Frontend/datos/ejemplo-panaderia.js"):
    wb = load_workbook(ruta, data_only=True)
    hojas = {}
    for ws in wb:
        if ws.title == "Verificacion":
            continue
        hojas[ws.title] = [[valor(c) for c in fila] for fila in ws.iter_rows(values_only=True)
                           if any(c is not None for c in fila)]
    with open(salida, "w", encoding="utf-8") as f:
        f.write("/* Datos de ejemplo: %s (simulación de Montecarlo, ver Simulacion/simular_datos.py).\n"
                " * Son datos SIMULADOS con autorización del docente, no mediciones reales. */\n" % ruta.split("/")[-1])
        f.write("window.DATOS_EJEMPLO = { nombre: \"Panadería San Nicolás · septiembre 2026 (simulado)\", hojas: ")
        f.write(json.dumps(hojas, ensure_ascii=False, separators=(",", ":")))
        f.write(" };\n")
    print("Escrito:", salida)


if __name__ == "__main__":
    main(*sys.argv[1:])
