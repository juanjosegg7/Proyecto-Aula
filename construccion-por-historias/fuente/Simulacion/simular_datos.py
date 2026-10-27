"""
Simulación de Montecarlo de un mes del proceso de pan trenza · Panadería San Nicolás.

Implementa los supuestos de Docs/Supuestos_Simulacion.md (documento
«Supuestos para la Simulación de Datos») y genera el libro de Excel con la
estructura de la plantilla Base_Panaderia_Indicadores.xlsx (tablas Dim y Fact).

Uso:
    python Simulacion/simular_datos.py                      # septiembre de 2026, semilla 2026
    python Simulacion/simular_datos.py --mes 10 --semilla 7 # otro mes para ensayar la demostración

Requiere Python 3.9+ y openpyxl (pip install openpyxl).
La semilla fija hace que los datos sean reproducibles (equivale a «pegar como valores» en Excel).
"""
import argparse
import calendar
import math
import random
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------- parámetros (sección 2 y 3)
PRODUCTO = {"producto_id": "PR_01", "nombre_producto": "Pan trenza 350 g", "tiempo_horneado_min": 22,
            "unidades_por_tanda": 16, "peso_g": 350}
TURNOS = [
    {"turno_id": "T1", "nombre_turno": "Turno 1", "hora_inicio": "06:00", "hora_fin": "14:00",
     "duracion_min": 480, "descanso_min": 30, "ventana_horno_min": 300},
    {"turno_id": "T2", "nombre_turno": "Turno 2", "hora_inicio": "14:00", "hora_fin": "22:00",
     "duracion_min": 480, "descanso_min": 30, "ventana_horno_min": 240},
]
MOTIVOS = [  # S = se excluye del tiempo programado (paro planeado)
    ("MP_01", "Precalentamiento del horno", "Planeado", "S"),
    ("MP_02", "Limpieza programada", "Planeado", "S"),
    ("MP_03", "Mantenimiento preventivo", "Planeado", "S"),
    ("MP_04", "Falta de masa (retraso de fermentación)", "No planeado", "N"),
    ("MP_05", "Falla del horno", "No planeado", "N"),
    ("MP_06", "Falta de gas o energía", "No planeado", "N"),
    ("MP_07", "Falta de personal", "No planeado", "N"),
    ("MP_08", "Horno asignado a otros productos", "Planeado", "S"),
]
OPERARIOS = {"T1": ["OP_01", "OP_02"], "T2": ["OP_03", "OP_04"]}
DEMANDA_MEDIA = {0: 170, 1: 160, 2: 165, 3: 175, 4: 210}  # lunes a viernes
CV_DEMANDA = 0.15
LAMBDA_PEDIDOS = 3
LAMBDA_PAROS = 1
P_DEFECTO = 0.04
COLCHON = 1.10
REPARTO_T1 = 0.60
MINUTOS_POR_TANDA_PLAN = 25  # 22 de horneado + 3 de carga y descarga
FESTIVOS = {date(2026, 10, 12), date(2026, 11, 2), date(2026, 11, 16), date(2026, 12, 8), date(2026, 12, 25)}


# ---------------------------------------------------------------- distribuciones (sección 5)
def normal_truncada(rng, media, desv, a, b):
    return min(b, max(a, rng.gauss(media, desv)))


def poisson(rng, lam):
    # Método de Knuth (adecuado para λ pequeño)
    limite, k, p = math.exp(-lam), 0, 1.0
    while True:
        p *= rng.random()
        if p <= limite:
            return k
        k += 1


def binomial(rng, n, p):
    return sum(1 for _ in range(n) if rng.random() < p)


def lognormal_media_desv(rng, m, s):
    mu = math.log(m * m / math.sqrt(m * m + s * s))
    sigma = math.sqrt(math.log(1 + s * s / (m * m)))
    return rng.lognormvariate(mu, sigma)


def paro_no_planeado(rng):
    u = rng.random()
    if u < 0.5:
        return "MP_04", rng.triangular(5, 25, 10)      # falta de masa: triangular(5, 10, 25)
    if u < 0.7:
        return "MP_05", lognormal_media_desv(rng, 30, 15)  # falla del horno
    if u < 0.9:
        return "MP_07", rng.triangular(10, 40, 20)     # falta de personal: triangular(10, 20, 40)
    return "MP_06", lognormal_media_desv(rng, 40, 20)  # falta de gas


# ---------------------------------------------------------------- simulación
def dias_habiles(anio, mes):
    _, n = calendar.monthrange(anio, mes)
    return [date(anio, mes, d) for d in range(1, n + 1)
            if date(anio, mes, d).weekday() < 5 and date(anio, mes, d) not in FESTIVOS]


def simular(anio, mes, semilla):
    rng = random.Random(semilla)
    prod, paros, pedidos, ventas, mano = [], [], [], [], []
    for dia in dias_habiles(anio, mes):
        media = DEMANDA_MEDIA[dia.weekday()]
        demanda_mostrador = max(0, round(rng.gauss(media, CV_DEMANDA * media)))

        # Pedidos del día
        pedidos_dia = []
        for _ in range(poisson(rng, LAMBDA_PEDIDOS)):
            cantidad = rng.randint(10, 40)
            compromiso = datetime.combine(dia, datetime.min.time()) + timedelta(hours=rng.randint(7, 18))
            desfase = rng.gauss(-15, 10)
            completo = rng.random() < 0.97
            faltante = 0 if completo else rng.randint(1, 5)
            pedidos_dia.append({"cantidad": cantidad, "compromiso": compromiso,
                                "entrega": compromiso + timedelta(minutes=round(desfase)),
                                "a_entregar": cantidad - faltante})
        planeada = round(media * COLCHON + sum(p["cantidad"] for p in pedidos_dia))
        plan_turno = {"T1": round(planeada * REPARTO_T1)}
        plan_turno["T2"] = planeada - plan_turno["T1"]

        buenas_dia = 0
        for t in TURNOS:
            tid = t["turno_id"]
            inicio = datetime.combine(dia, datetime.strptime(t["hora_inicio"], "%H:%M").time())
            reloj = inicio
            ventana = t["ventana_horno_min"]

            def registrar_paro(motivo, minutos, desde):
                paros.append({"fecha": dia, "turno_id": tid, "motivo_id": motivo,
                              "hora_inicio": desde, "hora_fin": desde + timedelta(minutes=minutos),
                              "duracion_min": round(minutos, 2)})
                return desde + timedelta(minutes=minutos)

            precal = rng.uniform(20, 30) if tid == "T1" else 0
            limpieza = rng.uniform(10, 20)
            no_planeados = [paro_no_planeado(rng) for _ in range(poisson(rng, LAMBDA_PAROS))]
            disponible = ventana - precal - limpieza - sum(m for _, m in no_planeados)
            necesarias = math.ceil(plan_turno[tid] / PRODUCTO["unidades_por_tanda"])
            caben = max(0, math.floor(disponible / MINUTOS_POR_TANDA_PLAN))
            tandas = min(necesarias, caben)

            if precal:
                reloj = registrar_paro("MP_01", precal, reloj)
            # Los paros no planeados se ubican entre tandas (regla 3)
            posiciones = sorted(rng.randint(0, max(tandas, 1)) for _ in no_planeados)
            k = 0
            for n in range(1, tandas + 1):
                while k < len(no_planeados) and posiciones[k] < n:
                    reloj = registrar_paro(no_planeados[k][0], no_planeados[k][1], reloj)
                    k += 1
                dur = normal_truncada(rng, 25, 1.5, 22, 30)
                defect = binomial(rng, PRODUCTO["unidades_por_tanda"], P_DEFECTO)
                prod.append({"fecha": dia, "turno_id": tid, "producto_id": PRODUCTO["producto_id"], "tanda": n,
                             "hora_inicio": reloj, "hora_fin": reloj + timedelta(minutes=dur),
                             "duracion_min": round(dur, 2), "unidades_producidas": PRODUCTO["unidades_por_tanda"],
                             "unidades_defectuosas": defect})
                reloj += timedelta(minutes=dur)
                buenas_dia += PRODUCTO["unidades_por_tanda"] - defect
            while k < len(no_planeados):
                reloj = registrar_paro(no_planeados[k][0], no_planeados[k][1], reloj)
                k += 1
            reloj = registrar_paro("MP_02", limpieza, reloj)
            # Resto del turno (menos el descanso): el horno hornea otros productos
            otros = t["duracion_min"] - t["descanso_min"] - ventana
            registrar_paro("MP_08", otros, reloj)

            # Mano de obra (regla 6: máximo 2 operarios y 8 h)
            for op in OPERARIOS[tid]:
                asistio = rng.random() < 0.97
                horas = round(normal_truncada(rng, 7.5, 0.25, 6, 8) * 4) / 4 if asistio else 0
                mano.append({"fecha": dia, "turno_id": tid, "operario_id": op,
                             "asistio": "S" if asistio else "N", "horas_trabajadas": horas})

        # Ventas (reglas 4 y 5): primero los pedidos, luego el mostrador
        disponible_venta = buenas_dia
        entregado_pedidos = 0
        for p in sorted(pedidos_dia, key=lambda x: x["compromiso"]):
            entregada = min(p["a_entregar"], disponible_venta)
            disponible_venta -= entregada
            entregado_pedidos += entregada
            pedidos.append({"fecha": dia, "cliente_codigo": "CL_%02d" % rng.randint(1, 15),
                            "producto_id": PRODUCTO["producto_id"], "cantidad_pedida": p["cantidad"],
                            "fecha_hora_compromiso": p["compromiso"], "fecha_hora_entrega": p["entrega"],
                            "cantidad_entregada": entregada})
        mostrador = min(demanda_mostrador, disponible_venta)
        ventas.append({"fecha": dia, "producto_id": PRODUCTO["producto_id"], "demanda_mostrador": demanda_mostrador,
                       "unidades_vendidas_mostrador": mostrador, "unidades_entregadas_pedidos": entregado_pedidos,
                       "unidades_vendidas": mostrador + entregado_pedidos})
    return prod, paros, pedidos, ventas, mano


# ---------------------------------------------------------------- indicadores (verificación, fichas técnicas)
def indicadores(prod, paros, pedidos, ventas, mano):
    excl = {m[0]: m[3] == "S" for m in MOTIVOS}
    turnos = {t["turno_id"]: t for t in TURNOS}
    claves = {(r["fecha"], r["turno_id"]) for r in prod} | {(r["fecha"], r["turno_id"]) for r in paros}
    programado = sum(turnos[t]["duracion_min"] - turnos[t]["descanso_min"] for _, t in claves) \
        - sum(p["duracion_min"] for p in paros if excl[p["motivo_id"]])
    no_plan = sum(p["duracion_min"] for p in paros if not excl[p["motivo_id"]])
    operativo = programado - no_plan
    ideal_u = PRODUCTO["tiempo_horneado_min"] / PRODUCTO["unidades_por_tanda"]
    unidades = sum(r["unidades_producidas"] for r in prod)
    defect = sum(r["unidades_defectuosas"] for r in prod)
    buenas = unidades - defect
    A, R, Q = operativo / programado, unidades * ideal_u / operativo, buenas / unidades
    sla = sum(1 for p in pedidos if p["fecha_hora_entrega"] <= p["fecha_hora_compromiso"]
              and p["cantidad_entregada"] >= p["cantidad_pedida"]) / len(pedidos)
    vendidas = sum(v["unidades_vendidas"] for v in ventas)
    horas = sum(m["horas_trabajadas"] for m in mano)
    return [("Disponibilidad", A), ("Rendimiento", R), ("Calidad", Q), ("OEE", A * R * Q), ("SLA (OTIF)", sla),
            ("% no vendido", (buenas - vendidas) / buenas), ("Productividad (trenzas/h-h)", buenas / horas),
            ("Tiempo programado (min)", programado), ("Tiempo operativo (min)", operativo),
            ("Unidades producidas", unidades), ("Unidades defectuosas", defect), ("Unidades vendidas", vendidas),
            ("Pedidos", len(pedidos)), ("Horas-hombre", horas)]


# ---------------------------------------------------------------- Excel
def hoja(wb, nombre, encabezados, filas, formatos=None):
    ws = wb.create_sheet(nombre)
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="9A3412")
    for f in filas:
        ws.append([f[h] if isinstance(f, dict) else f[i] for i, h in enumerate(encabezados)])
    for i, h in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 2)
        fmt = (formatos or {}).get(h)
        if fmt:
            for celda in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in celda:
                    c.number_format = fmt
    ws.freeze_panes = "A2"
    return ws


def con_ids(filas, campo):
    return [dict({campo: i + 1}, **f) for i, f in enumerate(filas)]


def escribir(ruta, prod, paros, pedidos, ventas, mano, anio, mes, semilla, vacio=False):
    wb = Workbook()
    wb.remove(wb.active)
    F, FH = "yyyy-mm-dd", "yyyy-mm-dd hh:mm"
    hoja(wb, "Dim_Producto", list(PRODUCTO), [PRODUCTO])
    hoja(wb, "Dim_Turno", ["turno_id", "nombre_turno", "hora_inicio", "hora_fin", "duracion_min", "descanso_min"], TURNOS)
    hoja(wb, "Dim_MotivoParo", ["motivo_id", "nombre_motivo", "tipo", "excluir_tiempo_programado"], MOTIVOS)
    hoja(wb, "Dim_Operario", ["operario_id", "nombre_operario", "turno_habitual"],
         [(op, "Operario %s" % op[-1], t) for t, ops in OPERARIOS.items() for op in ops])
    hoja(wb, "Fact_Produccion", ["produccion_id", "fecha", "turno_id", "producto_id", "tanda", "hora_inicio", "hora_fin",
                                 "duracion_min", "unidades_producidas", "unidades_defectuosas"],
         [] if vacio else con_ids(prod, "produccion_id"), {"fecha": F, "hora_inicio": FH, "hora_fin": FH})
    hoja(wb, "Fact_Paros", ["paro_id", "fecha", "turno_id", "motivo_id", "hora_inicio", "hora_fin", "duracion_min"],
         [] if vacio else con_ids(paros, "paro_id"), {"fecha": F, "hora_inicio": FH, "hora_fin": FH})
    hoja(wb, "Fact_Pedidos", ["pedido_id", "fecha", "cliente_codigo", "producto_id", "cantidad_pedida",
                              "fecha_hora_compromiso", "fecha_hora_entrega", "cantidad_entregada"],
         [] if vacio else con_ids(pedidos, "pedido_id"),
         {"fecha": F, "fecha_hora_compromiso": FH, "fecha_hora_entrega": FH})
    hoja(wb, "Fact_Ventas", ["fecha", "producto_id", "demanda_mostrador", "unidades_vendidas_mostrador",
                             "unidades_entregadas_pedidos", "unidades_vendidas"], [] if vacio else ventas, {"fecha": F})
    hoja(wb, "Fact_ManoObra", ["registro_id", "fecha", "turno_id", "operario_id", "asistio", "horas_trabajadas"],
         [] if vacio else con_ids(mano, "registro_id"), {"fecha": F})
    if not vacio:
        ws = hoja(wb, "Verificacion", ["Indicador", "Valor"], indicadores(prod, paros, pedidos, ventas, mano))
        ws.append([])
        ws.append(["Datos simulados con autorización del docente", "Semilla %d · %02d/%d" % (semilla, mes, anio)])
        ws.append(["Supuestos", "Docs/Supuestos_Simulacion.md"])
    wb.save(ruta)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--anio", type=int, default=2026)
    ap.add_argument("--mes", type=int, default=9)
    ap.add_argument("--semilla", type=int, default=2026)
    ap.add_argument("--salida", default=None)
    ap.add_argument("--plantilla", default=None, help="también escribe la plantilla vacía en esta ruta")
    a = ap.parse_args()
    datos = simular(a.anio, a.mes, a.semilla)
    salida = a.salida or "Base_Panaderia_Indicadores_%d-%02d.xlsx" % (a.anio, a.mes)
    escribir(salida, *datos, a.anio, a.mes, a.semilla)
    if a.plantilla:
        escribir(a.plantilla, *datos, a.anio, a.mes, a.semilla, vacio=True)
    print("Archivo:", salida)
    for nombre, valor in indicadores(*datos):
        print("  %-28s %s" % (nombre, ("%.2f %%" % (valor * 100)) if valor <= 1.5 else ("%.2f" % valor)))
